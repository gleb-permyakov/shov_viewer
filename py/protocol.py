from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class ProtocolTemplateError(RuntimeError):
    pass


def _set_paragraph_run_text(paragraph, run_index: int, value: Any) -> None:
    try:
        paragraph.runs[run_index].text = "" if value is None else str(value)
    except IndexError as exc:
        raise ProtocolTemplateError(
            f"В шаблоне изменилась структура абзаца; не найден run {run_index}."
        ) from exc


def _set_cell_text_preserve_style(cell, value: Any) -> None:
    text = "" if value is None else str(value)

    if not cell.paragraphs:
        cell.text = text
        return

    paragraph = cell.paragraphs[0]

    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _replace_static_placeholders(doc: Document, data: dict[str, Any]) -> None:
    paragraphs = doc.paragraphs

    # код просмотра структуры документа
    # counter = 0
    # for par in paragraphs:
    #     print("\n======= PARAGRAPH " + str(counter))
    #     counter_2 = 0
    #     for j in par.runs:
    #         print("==> RUN " + str(counter_2) + ")'" + j.text + "'")
    #         counter_2 += 1
    #     counter += 1
        
    # Заголовок протокола
    p = paragraphs[1]
    _set_paragraph_run_text(p, 1, "№ " + "1" + " ")
    _set_paragraph_run_text(p, 2, "от " + "25.04.2026" + " г.")
    _set_paragraph_run_text(p, 3, "")
    _set_paragraph_run_text(p, 4, "\nпроверки сварных стыков " + data.get("object_name", "") + " радиографическим методом" )

    # Адрес объекта
    p = paragraphs[2]
    _set_paragraph_run_text(p, 1, "сварных стыков "+ data.get("object_name_2", "") + ", стоящего по адресу:")
    _set_paragraph_run_text(p, 2, "\n" + data.get("address", ""))

    # Параметры трубы
    p = paragraphs[3]
    _set_paragraph_run_text(p, 0, data.get("made_of", ""))
    _set_paragraph_run_text(p, 1, "")
    _set_paragraph_run_text(p, 2, "")

    # Заключение лаборатории
    conclusion = data.get("lab_conclusion")
    if conclusion:
        p = paragraphs[6]
        _set_paragraph_run_text(p, 0, "Заключение лаборатории: " + conclusion)

    # Подпись начальника лабаратории
    p = paragraphs[7]
    _set_paragraph_run_text(p, 0, "Начальник лаборатории " + data.get("chief_name", ""))

    # Подпись дефектоскописта
    p = paragraphs[8]
    _set_paragraph_run_text(p, 0, "Дефектоскопист " + data.get("inspector_name", ""))
    _set_paragraph_run_text(p, 11, "    " + "25.04.2026")


def _replace_table_header_placeholders(table, data: dict[str, Any]) -> None:
    header_cell = table.rows[0].cells[7]
    runs = header_cell.paragraphs[0].runs


def _build_defects_table(doc: Document, defects: list[dict[str, Any]], data: dict[str, Any]) -> None:
    if not doc.tables:
        raise ProtocolTemplateError("В шаблоне не найдена таблица результатов проверки.")

    table = doc.tables[0]
    _replace_table_header_placeholders(table, data)

    if len(table.rows) < 2:
        raise ProtocolTemplateError("В шаблоне нет образца строки таблицы для копирования.")

    sample_row_xml = deepcopy(table.rows[1]._tr)

    # Удаляем все строки кроме заголовка
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[1]._tr)

    # Генерируем строки по дефектам
    for defect in defects:
        new_row_xml = deepcopy(sample_row_xml)
        table._tbl.append(new_row_xml)

        row = table.rows[-1]
        values = [ "", "", "", "", "", defect, "", "", "", "", "", ]

        for cell, value in zip(row.cells, values):
            _set_cell_text_preserve_style(cell, value)


def build_protocol_doc(data: dict[str, Any], template_path: str | Path) -> BytesIO:
    doc = Document(str(template_path))
    
    _replace_static_placeholders(doc, data)
    _build_defects_table(doc, data.get("defects", []), data)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
